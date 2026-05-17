"""Build the recruiters XLSX from docs/recruiters.csv.

The CSV is the source of truth. This script just renders an Excel version
with clickable URLs and bold headers for people who prefer a spreadsheet
app over a CSV.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
CSV_PATH = ROOT / "docs" / "recruiters.csv"
XLSX_PATH = ROOT / "docs" / "recruiters.xlsx"


def main() -> None:
    rows = []
    with CSV_PATH.open(encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            rows.append((r["Company"], r["Type"], r["URL"]))

    wb = Workbook()
    ws = wb.active
    ws.title = "Recruiters"
    ws.append(["Company", "Type", "URL"])
    for cell in ws[1]:
        cell.font = Font(bold=True)

    link_font = Font(color="0563C1", underline="single")
    for company, type_, url in rows:
        ws.append([company, type_, url])
        if url:
            cell = ws.cell(row=ws.max_row, column=3)
            cell.hyperlink = url
            cell.font = link_font

    for col, width in {1: 38, 2: 16, 3: 62}.items():
        ws.column_dimensions[get_column_letter(col)].width = width

    wb.save(XLSX_PATH)

    counts: dict[str, int] = {}
    for _, t, _ in rows:
        counts[t] = counts.get(t, 0) + 1
    print(f"Wrote {XLSX_PATH} ({len(rows)} rows)")
    for t in sorted(counts):
        print(f"  {t:18s} {counts[t]:3d}")


if __name__ == "__main__":
    main()
